# -*- coding: utf-8 -*-
"""Генератор чек-листа монетизации (.xlsx) для владельца игры Mining.
Собирает все gamepass / devproduct / group / universe / иконки, которым нужен
ручной ID из Creator Hub, и раскладывает по колонкам с путями к картинкам.
"""
import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

ROOT = r"C:\Projects\Roblox\Mining"
ASSETS = os.path.join(ROOT, "assets", "ui")
OUT = os.path.join(ROOT, "Монетизация_ID_и_картинки.xlsx")


def img(name):
    """Абсолютный путь к png в assets/ui, если файл существует."""
    p = os.path.join(ASSETS, name + ".png")
    return p if os.path.exists(p) else None


HEADERS = [
    "Ключ", "Название (RU)", "Название (EN)", "Тип", "Цена (R$)",
    "Текущий ID", "Что нужно сделать", "Файл изображения (путь)",
    "Статус картинки", "Куда вписать ID", "Примечание",
]

# type_order: gamepass=0, devproduct=1, group/universe=2, icon=3
rows = []


def add(key, ru, en, typ, price, cur_id, todo, icon_key, img_status,
        where, note, order):
    rows.append({
        "key": key, "ru": ru, "en": en, "typ": typ, "price": price,
        "cur_id": cur_id, "todo": todo, "icon": icon_key,
        "img_status": img_status, "where": where, "note": note, "order": order,
    })


# ---------- GAMEPASSES ----------
add("vip", "VIP", "VIP", "Game Pass", 399, "0 (плейсхолдер)",
    "Создать Game Pass в Creator Hub, вставить ID в constants.lua",
    "icon_crown", None,
    "constants.lua → GAMEPASSES.vip.id",
    "+10% монет, титул VIP и золотой ник (+2 слота питомцев: base 3 → 5)", 0)
add("autoSell", "Auto-Sell", "Auto-Sell", "Game Pass", 599, "0 (плейсхолдер)",
    "Создать Game Pass, вставить ID в constants.lua",
    "upg_autosell", None,
    "constants.lua → GAMEPASSES.autoSell.id",
    "Авто-продажа навсегда — инвентарь не переполнится. Бейдж ПОПУЛЯРНО", 0)
add("petSlots", "+2 слота питомцев", "+2 Pet Slots", "Game Pass", 799,
    "0 (плейсхолдер)",
    "Создать Game Pass, вставить ID в constants.lua",
    "tab_pets", None,
    "constants.lua → GAMEPASSES.petSlots.id",
    "Legacy: слоты питомцев теперь считаются моделью PETS.slots. Пасс можно "
    "перепрофилировать или убрать (см. отчёт).", 0)

# ---------- DEVPRODUCTS ----------
DP = [
    ("starterPack", "Стартовый набор", "Starter Pack", 99, "icon: pack_starter",
     "Набор (bundle): 25 000 монет + x2 монеты 30 мин + 3 яйца. Старая цена 499, бейдж СТАРТ. oneTime.", "pack_starter"),
    ("bundleMiner", "Набор шахтёра", "Miner Bundle", 349, "",
     "Набор: 50 000 монет + x2 удача 15 мин + 5 яиц. Старая цена 699, бейдж -50%.", "pack_miner"),
    ("bundleMega", "Мега набор", "Mega Bundle", 799, "",
     "Набор: 250 000 монет + x2 монеты 1ч + x2 сила 30 мин + 15 яиц. Старая цена 1599, бейдж ХИТ.", "pack_mega"),
    ("boostLuck15", "Удача x2", "Luck x2", 49, "",
     "Буст удачи x2 на 15 мин (больше редкой руды и комнат).", "buff_luck"),
    ("boostLuck60", "Удача x2", "Luck x2", 149, "",
     "Буст удачи x2 на 1 час. Старая цена 196.", "buff_luck"),
    ("boostCoins15", "Монеты x2", "Coins x2", 49, "",
     "Буст монет x2 на 15 мин (удвоенная продажа).", "buff_coin"),
    ("boostCoins60", "Монеты x2", "Coins x2", 149, "",
     "Буст монет x2 на 1 час. Старая цена 196.", "buff_coin"),
    ("boostDamage15", "Сила x2", "Power x2", 59, "",
     "Буст урона x2 на 15 мин.", "buff_damage"),
    ("boostSpeed15", "Скорость x2", "Speed x2", 59, "",
     "Буст скорости копания x2 на 15 мин.", "upg_speed"),
    ("coinsSmall", "Малый пакет", "Small Pack", 99, "",
     "+10 000 монет. Старая цена 149.", "coin"),
    ("coinsMedium", "Средний пакет", "Medium Pack", 399, "",
     "+100 000 монет. Старая цена 599.", "coin"),
    ("coinsLarge", "Большой пакет", "Large Pack", 899, "",
     "+500 000 монет. Старая цена 1499, бейдж -40%.", "coin"),
    ("coinsMega", "Мега пакет", "Mega Pack", 2499, "",
     "+2 000 000 монет. Старая цена 4999, бейдж ХИТ.", "coin"),
    ("egg5", "5 яиц", "5 Eggs", 99, "",
     "5 вылуплений подряд.", "icon_egg"),
    ("egg10", "10 яиц", "10 Eggs", 199, "",
     "10 яиц сразу. Старая цена 249.", "icon_egg"),
    ("egg25", "25 яиц", "25 Eggs", 399, "",
     "25 яиц — шанс на легенду. Старая цена 598, бейдж -33%.", "icon_egg"),
]
for key, ru, en, price, _extra, note, icon_key in DP:
    add(key, ru, en, "Dev Product", price, "0 (плейсхолдер)",
        "Создать Developer Product, вставить ID в constants.lua",
        icon_key, None,
        f"constants.lua → DEVPRODUCTS.{key}.id",
        note, 1)

# ---------- GROUP / UNIVERSE ----------
add("SOCIAL_REWARD.groupId", "ID группы (соц-награда)", "Group ID",
    "Group", "", "0 (плейсхолдер)",
    "Создать/взять группу в Creator Hub, вставить её ID",
    "icon_social_reward", "уже загружено (icon_social_reward=126058376087998)",
    "constants.lua → SOCIAL_REWARD.groupId",
    "Награда за вступление в группу: 7500 монет + 15 кристаллов + x2 монеты 15 мин.", 2)
add("SOCIAL_REWARD.universeId", "ID вселенной (universe)", "Universe ID",
    "Universe", "", "0 (плейсхолдер)",
    "Взять Universe ID игры (НЕ placeId!) из Creator Hub, вставить в constants.lua",
    None, None,
    "constants.lua → SOCIAL_REWARD.universeId",
    "Нужен для соц-награды «добавить игру в избранное».", 2)

# ---------- ICONS (уже загружены — справочно) ----------
add("loading_bg", "Фон загрузочного экрана", "Loading background",
    "Иконка", "", "109016788278844",
    "Ничего — уже загружено", "loading_bg", "уже загружено",
    "UiAssets.lua → ROBLOX_IMAGES.loading_bg",
    "Фон LoadingScreen. Загружен.", 3)
add("icon_social_reward", "Иконка соц-награды", "Social reward icon",
    "Иконка", "", "126058376087998",
    "Ничего — уже загружено", "icon_social_reward", "уже загружено",
    "UiAssets.lua → ROBLOX_IMAGES.icon_social_reward",
    "Загружен.", 3)
add("icon_promo_code", "Иконка промокода", "Promo code icon",
    "Иконка", "", "105183216335516",
    "Ничего — уже загружено", "icon_promo_code", "уже загружено",
    "UiAssets.lua → ROBLOX_IMAGES.icon_promo_code",
    "Загружен.", 3)

rows.sort(key=lambda r: (r["order"], r["key"].lower()))

# ---------- BUILD WORKBOOK ----------
wb = Workbook()
ws = wb.active
ws.title = "Монетизация"

header_fill = PatternFill("solid", fgColor="2F5496")
header_font = Font(bold=True, color="FFFFFF", size=11)
thin = Side(style="thin", color="D0D0D0")
border = Border(left=thin, right=thin, top=thin, bottom=thin)
wrap = Alignment(vertical="top", wrap_text=True)

for c, h in enumerate(HEADERS, 1):
    cell = ws.cell(row=1, column=c, value=h)
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(vertical="center", horizontal="center", wrap_text=True)
    cell.border = border

link_font = Font(color="0563C1", underline="single")

r = 2
for row in rows:
    path = img(row["icon"]) if row["icon"] else None
    if row["icon"] and path:
        img_status = row["img_status"] or "есть"
    elif row["icon"]:
        img_status = "нет изображения — нужно создать"
    else:
        img_status = row["img_status"] or "нет изображения — нужно создать"

    values = [
        row["key"], row["ru"], row["en"], row["typ"],
        row["price"] if row["price"] != "" else "—",
        row["cur_id"], row["todo"],
        path if path else ("нет изображения — нужно создать" if not row["icon"] else "нет файла"),
        img_status, row["where"], row["note"],
    ]
    for c, v in enumerate(values, 1):
        cell = ws.cell(row=r, column=c, value=v)
        cell.alignment = wrap
        cell.border = border
    # hyperlink for image path
    if path:
        pc = ws.cell(row=r, column=8)
        pc.hyperlink = "file:///" + path.replace("\\", "/")
        pc.font = link_font
    r += 1

# column widths
widths = [16, 22, 18, 13, 9, 18, 40, 46, 32, 40, 50]
for i, w in enumerate(widths, 1):
    ws.column_dimensions[get_column_letter(i)].width = w

ws.freeze_panes = "A2"
ws.auto_filter.ref = f"A1:{get_column_letter(len(HEADERS))}{r-1}"
ws.row_dimensions[1].height = 32

wb.save(OUT)
print("SAVED:", OUT)
print("ROWS:", r - 2)
