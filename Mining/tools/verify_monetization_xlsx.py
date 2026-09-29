# -*- coding: utf-8 -*-
import os
from openpyxl import load_workbook

OUT = r"C:\Projects\Roblox\Mining\Монетизация_ID_и_картинки.xlsx"
print("EXISTS:", os.path.exists(OUT), "SIZE:", os.path.getsize(OUT), "bytes")
wb = load_workbook(OUT)
ws = wb.active
print("SHEET:", ws.title, "| DIM:", ws.dimensions, "| DATA ROWS:", ws.max_row - 1)
print("FREEZE:", ws.freeze_panes, "| FILTER:", ws.auto_filter.ref)
print("-" * 90)
missing_img = []
for row in ws.iter_rows(min_row=2, values_only=False):
    key = row[0].value
    typ = row[3].value
    price = row[4].value
    cur = row[5].value
    imgstat = row[8].value
    link = row[7].hyperlink.target if row[7].hyperlink else None
    print(f"{key:26} | {typ:12} | R$ {str(price):6} | id={str(cur):16} | {imgstat}")
    if "нужно создать" in str(imgstat):
        missing_img.append(key)
print("-" * 90)
print("MISSING IMAGES:", missing_img if missing_img else "none")
